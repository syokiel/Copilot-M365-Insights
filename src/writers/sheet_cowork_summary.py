from collections import defaultdict
from datetime import datetime, timezone

from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from src.writers._style import HEADER_FILL, HEADER_FONT, LEFT


def write(
    ws: Worksheet,
    cowork_usage: list[dict],
    consumption_by_service: list[dict] | None = None,
) -> None:
    consumption_by_service = consumption_by_service or []
    cowork_credits = [r for r in consumption_by_service if (r.get("service_name") or "") == "Cowork"]

    # ── Aggregates: task activity ────────────────────────────────────────────
    total_users      = len(cowork_usage)
    total_tasks      = sum(r.get("total_tasks") or 0 for r in cowork_usage)
    scheduled_tasks  = sum(r.get("scheduled_tasks") or 0 for r in cowork_usage)
    initiated_tasks  = sum(r.get("user_initiated_tasks") or 0 for r in cowork_usage)
    total_active_days = sum(r.get("active_days") or 0 for r in cowork_usage)
    avg_tasks_per_user = round(total_tasks / total_users, 1) if total_users else 0.0
    avg_active_days    = round(total_active_days / total_users, 1) if total_users else 0.0

    top_users = sorted(cowork_usage, key=lambda r: r.get("total_tasks") or 0, reverse=True)[:10]

    # ── Aggregates: credit consumption (Viva Consumption, service = "Cowork") ─
    total_credits   = sum(r.get("total_credits_used") or 0 for r in cowork_credits)
    people_consuming = {r.get("people_historical_id") for r in cowork_credits if r.get("people_historical_id")}
    avg_credits_per_person = round(total_credits / len(people_consuming), 1) if people_consuming else 0.0

    org_credits: dict[str, float] = defaultdict(float)
    for r in cowork_credits:
        org_credits[r.get("organization") or "Unknown"] += r.get("total_credits_used") or 0
    top_orgs = sorted(org_credits.items(), key=lambda x: x[1], reverse=True)[:10]

    def _fmt(v) -> str:
        if v is None:
            return "—"
        if isinstance(v, float):
            return f"{v:,.1f}"
        return str(v)

    rows: list[tuple] = [
        ("Report generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        (None, None),
        ("── Task Activity (M365ADMIN_COWORK_USAGE) ────────", None),
        ("Total Cowork users",         total_users),
        ("Total tasks",                total_tasks),
        ("Scheduled tasks",            scheduled_tasks),
        ("User-initiated tasks",       initiated_tasks),
        ("Avg tasks per user",         avg_tasks_per_user),
        ("Avg active days per user",   avg_active_days),
    ]

    if top_users:
        rows += [
            (None, None),
            ("── Top Users by Total Tasks ──────────────────", None),
        ]
        for r in top_users:
            label = r.get("display_name") or r.get("user_principal_name") or "Unknown"
            rows.append((f"  {label}", r.get("total_tasks")))

    rows += [
        (None, None),
        ("── Credit Consumption (VIVA_REPORT_CONSUMPTION) ──", None),
        ("Total Cowork credits used",      _fmt(total_credits)),
        ("People consuming Cowork credits", len(people_consuming)),
        ("Avg credits per person",         _fmt(avg_credits_per_person)),
    ]

    if top_orgs:
        rows += [
            (None, None),
            ("── Top Organizations by Cowork Credits ───────", None),
        ]
        for org, total in top_orgs:
            rows.append((f"  {org}", _fmt(total)))

    if not cowork_usage and not cowork_credits:
        rows.append((None, None))
        rows.append(("  No Cowork data imported (set M365ADMIN_COWORK_USAGE / VIVA_REPORT_CONSUMPTION)", None))

    for col, header in enumerate(["Metric", "Value"], 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = LEFT

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 28

    for row_idx, (metric, value) in enumerate(rows, 2):
        a = ws.cell(row=row_idx, column=1, value=metric)
        b = ws.cell(row=row_idx, column=2, value=value)
        a.alignment = LEFT
        b.alignment = LEFT
        if metric and metric.startswith("──"):
            a.font = Font(bold=True, size=11, color="1F4E79")
        else:
            a.font = Font(size=11)
            b.font = Font(size=11)
