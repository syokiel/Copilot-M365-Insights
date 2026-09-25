"""DailyCredit_Trends — daily and weekly credit consumption by service from the
Viva daily consumption export (M365 services such as Cowork, and GitHub AI),
with 7-day rolling averages and charts.

Within a service's coverage (its first to last date in the export), days
with no rows count as 0 credits, so averages reflect the calendar rather than
only active days. Outside that coverage cells are blank, not 0."""
from datetime import date, timedelta

from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.writers._style import apply_row_style, write_headers

GITHUB = "GitHub AI"
_TITLE_FONT = Font(bold=True, size=12, color="1F4E79")
_NOTE_FONT = Font(italic=True, size=9, color="595959")
_CREDIT_FMT = "#,##0.0"
_COUNT_FMT = "#,##0"
_CHART_ROWS = 16


def _services(rows: list[dict]) -> list[str]:
    """M365 services alphabetically, GitHub AI last (it's on its own scale)."""
    names = {r["service"] for r in rows}
    return sorted(n for n in names if n != GITHUB) + ([GITHUB] if GITHUB in names else [])


def _daily_matrix(rows: list[dict], services: list[str]) -> list[dict]:
    by_key = {(r["metric_date"], r["service"]): r for r in rows}
    coverage = {
        s: (min(r["metric_date"] for r in rows if r["service"] == s),
            max(r["metric_date"] for r in rows if r["service"] == s))
        for s in services
    }
    first = date.fromisoformat(min(c[0] for c in coverage.values()))
    last = date.fromisoformat(max(c[1] for c in coverage.values()))
    out, window, gh_window = [], [], []
    d = first
    while d <= last:
        iso = d.isoformat()
        rec = {"date": iso}
        total, covered = 0.0, False
        for s in services:
            if not coverage[s][0] <= iso <= coverage[s][1]:
                continue   # outside this service's export window: leave blank
            r = by_key.get((iso, s)) or {}
            credits = r.get("credits") or 0.0
            rec[f"{s} credits"] = credits
            rec[f"{s} people"] = r.get("people") or 0
            total += credits
            covered = True
        if covered:
            rec["total"] = total
            window = (window + [total])[-7:]
            rec["avg7"] = sum(window) / len(window)
        if f"{GITHUB} credits" in rec:
            gh_window = (gh_window + [rec[f"{GITHUB} credits"]])[-7:]
            rec["gh_avg7"] = sum(gh_window) / len(gh_window)
        out.append(rec)
        d += timedelta(days=1)
    return out


def _add_chart(ws, title, header_row, last_row, cols, cat_col, anchor, y_title="Credits"):
    if not cols:
        return
    chart = LineChart()
    chart.title = title
    chart.y_axis.title = y_title
    chart.height, chart.width = 7.5, 18
    for c in cols:
        chart.add_data(Reference(ws, min_col=c, min_row=header_row, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=cat_col, min_row=header_row + 1, max_row=last_row))
    ws.add_chart(chart, anchor)


def write(ws: Worksheet, daily_totals: list[dict], weekly_totals: list[dict]) -> None:
    if not daily_totals:
        ws.cell(row=1, column=1,
                value="— No daily consumption data imported (set VIVA_REPORT_CONSUMPTION_DAILY) —")
        return

    services = _services(daily_totals)
    m365 = [s for s in services if s != GITHUB]

    # ── Daily table ─────────────────────────────────────────────────────────
    ws.cell(row=1, column=1, value="Daily credits by service").font = _TITLE_FONT
    ws.cell(row=2, column=1, value="Missing days count as 0. 7-day averages are trailing, per calendar day."
            ).font = _NOTE_FONT
    columns = [("Date", "date", None)]
    columns += [(f"{s} Credits", f"{s} credits", _CREDIT_FMT) for s in services]
    columns += [("Total Credits", "total", _CREDIT_FMT), ("Total 7-day Avg", "avg7", _CREDIT_FMT)]
    if GITHUB in services:
        columns += [(f"{GITHUB} 7-day Avg", "gh_avg7", _CREDIT_FMT)]
    columns += [(f"{s} People", f"{s} people", _COUNT_FMT) for s in services]

    header_row = 3
    write_headers(ws, [c[0] for c in columns], start_row=header_row)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    matrix = _daily_matrix(daily_totals, services)
    for i, rec in enumerate(matrix, start=header_row + 1):
        for j, (_h, key, fmt) in enumerate(columns, start=1):
            cell = ws.cell(row=i, column=j, value=rec.get(key))
            if fmt:
                cell.number_format = fmt
        apply_row_style(ws, i, len(columns))
    last_row = header_row + len(matrix)
    col_of = {key: j for j, (_h, key, _f) in enumerate(columns, start=1)}
    for j, (header, _k, _f) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, len(header) + 3)

    # Charts to the right of the daily table.
    chart_col = get_column_letter(len(columns) + 2)
    anchor_row = header_row
    if len(matrix) >= 2:
        if m365:
            _add_chart(ws, "M365 service credits per day", header_row, last_row,
                       [col_of[f"{s} credits"] for s in m365], 1, f"{chart_col}{anchor_row}")
            anchor_row += _CHART_ROWS
        if GITHUB in services:
            _add_chart(ws, "GitHub AI credits per day (with 7-day average)", header_row, last_row,
                       [col_of[f"{GITHUB} credits"], col_of["gh_avg7"]], 1, f"{chart_col}{anchor_row}")
            anchor_row += _CHART_ROWS
        _add_chart(ws, "People consuming credits per day", header_row, last_row,
                   [col_of[f"{s} people"] for s in services], 1, f"{chart_col}{anchor_row}", y_title="People")
        anchor_row += _CHART_ROWS

    # ── Weekly table ────────────────────────────────────────────────────────
    if not weekly_totals:
        return
    start = max(last_row, anchor_row) + 3
    ws.cell(row=start, column=1, value="Weekly credits by service (weeks start Monday)").font = _TITLE_FONT
    ws.cell(row=start + 1, column=1,
            value="People are distinct within the week. A partial first/last week shows fewer days."
            ).font = _NOTE_FONT
    weeks = sorted({r["week"] for r in weekly_totals})
    by_key = {(r["week"], r["service"]): r for r in weekly_totals}
    wcols = ["Week Starting"] + [f"{s} Credits" for s in services] + ["Total Credits"] \
        + [f"{s} People" for s in services] + [f"{s} Days" for s in services]
    wheader = start + 2
    write_headers(ws, wcols, start_row=wheader)
    for i, week in enumerate(weeks, start=wheader + 1):
        credits = [(by_key.get((week, s)) or {}).get("credits") or 0.0 for s in services]
        people = [(by_key.get((week, s)) or {}).get("people") or 0 for s in services]
        days = [(by_key.get((week, s)) or {}).get("days") or 0 for s in services]
        values = [week] + credits + [sum(credits)] + people + days
        for j, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            if 2 <= j <= len(services) + 2:
                cell.number_format = _CREDIT_FMT
            elif j > 1:
                cell.number_format = _COUNT_FMT
        apply_row_style(ws, i, len(wcols))
    wlast = wheader + len(weeks)
    if len(weeks) >= 2:
        anchor = wheader
        if m365:
            _add_chart(ws, "M365 service credits per week", wheader, wlast,
                       [2 + services.index(s) for s in m365], 1, f"{chart_col}{anchor}")
            anchor += _CHART_ROWS
        if GITHUB in services:
            _add_chart(ws, "GitHub AI credits per week", wheader, wlast,
                       [2 + services.index(GITHUB)], 1, f"{chart_col}{anchor}")
