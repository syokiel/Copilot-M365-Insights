"""Credits_Daily — daily credit consumption from the Viva daily consumption
export: totals per service per day (M365 services such as Cowork, plus
GitHub AI), then per-person detail for the M365 services."""
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.writers._style import apply_row_style, autofit_columns, write_headers

TOTAL_HEADERS = ["Date", "Service", "People", "Sessions", "Credits"]
_TOTAL_FIELDS = ["metric_date", "service", "people", "sessions", "credits"]

DETAIL_HEADERS = [
    "Date", "Person ID", "Service", "Sessions", "Credits", "User Limit",
    "Spending Policy", "Organization", "Copilot Licensed",
]
_DETAIL_FIELDS = [
    "metric_date", "person_id", "service_name", "session_count", "total_credits_used",
    "user_limit", "spending_policy", "organization", "is_copilot_licensed",
]


def write(ws: Worksheet, totals: list[dict], detail: list[dict]) -> None:
    write_headers(ws, TOTAL_HEADERS)
    if not totals:
        ws.cell(row=2, column=1, value="— No daily consumption data imported (set VIVA_REPORT_CONSUMPTION_DAILY) —")
        autofit_columns(ws, TOTAL_HEADERS)
        return
    for i, r in enumerate(totals, start=2):
        for col, field in enumerate(_TOTAL_FIELDS, start=1):
            cell = ws.cell(row=i, column=col, value=r.get(field))
            if field == "credits":
                cell.number_format = "#,##0.00"
        apply_row_style(ws, i, len(TOTAL_HEADERS))
    autofit_columns(ws, TOTAL_HEADERS)

    if not detail:
        return
    start = len(totals) + 4
    ws.cell(row=start - 1, column=1, value="Per-person daily detail (M365 services)").font = Font(
        bold=True, size=12, color="1F4E79")
    write_headers(ws, DETAIL_HEADERS, start_row=start)
    for i, r in enumerate(detail, start=start + 1):
        for col, field in enumerate(_DETAIL_FIELDS, start=1):
            value = r.get(field)
            if field == "is_copilot_licensed" and value is not None:
                value = "Yes" if value else "No"
            cell = ws.cell(row=i, column=col, value=value)
            if field == "total_credits_used":
                cell.number_format = "#,##0.00"
        apply_row_style(ws, i, len(DETAIL_HEADERS))
    # The detail table sits under the totals; widen shared columns to fit it.
    for col, header in enumerate(DETAIL_HEADERS, start=1):
        dim = ws.column_dimensions[get_column_letter(col)]
        dim.width = max(dim.width or 0, 40 if header == "Person ID" else len(header) + 4)
