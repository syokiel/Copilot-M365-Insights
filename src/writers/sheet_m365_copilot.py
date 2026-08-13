from openpyxl.worksheet.worksheet import Worksheet

from src.writers._style import apply_row_style, autofit_columns, write_headers

HEADERS = [
    "User Principal Name",
    "Display Name",
    "Last Activity Date",
    "Teams Chats",
    "Teams Meetings",
    "Word",
    "Excel",
    "PowerPoint",
    "Outlook",
    "OneNote",
    "Loop",
    "Copilot Chat",
    "Report Period",
    "Report Refresh Date",
    # ── Optional columns from the M365ADMIN_USAGE_COPILOT CSV export ────────
    "Prompts (All Apps)",
    "Prompts (Copilot Chat Work)",
    "Prompts (Copilot Chat Web)",
    "Active Usage Days (All Apps)",
    "Last Activity — Copilot Chat (Work)",
    "Last Activity — Copilot Chat (Web)",
    "Last Activity — Teams Copilot",
    "Last Activity — Word Copilot",
    "Last Activity — Excel Copilot",
    "Last Activity — PowerPoint Copilot",
    "Last Activity — Outlook Copilot",
    "Last Activity — OneNote Copilot",
    "Last Activity — Loop Copilot",
    "Last Activity — M365 Copilot App",
    "Last Activity — Edge",
]


def write(ws: Worksheet, rows: list[dict]) -> None:
    write_headers(ws, HEADERS)

    if not rows:
        ws.cell(row=2, column=1, value="— Requires Reports.Read.All permission on the sync service principal, or set M365ADMIN_USAGE_COPILOT —")
    else:
        for i, r in enumerate(rows, start=2):
            ws.cell(row=i, column=1, value=r.get("user_principal_name", ""))
            ws.cell(row=i, column=2, value=r.get("display_name", ""))
            ws.cell(row=i, column=3, value=r.get("last_activity_date", ""))
            ws.cell(row=i, column=4, value=r.get("teams_chats"))
            ws.cell(row=i, column=5, value=r.get("teams_meetings"))
            ws.cell(row=i, column=6, value=r.get("word"))
            ws.cell(row=i, column=7, value=r.get("excel"))
            ws.cell(row=i, column=8, value=r.get("powerpoint"))
            ws.cell(row=i, column=9, value=r.get("outlook"))
            ws.cell(row=i, column=10, value=r.get("onenote"))
            ws.cell(row=i, column=11, value=r.get("loop"))
            ws.cell(row=i, column=12, value=r.get("copilot_chat"))
            ws.cell(row=i, column=13, value=r.get("report_period", ""))
            ws.cell(row=i, column=14, value=r.get("report_refresh_date", ""))
            ws.cell(row=i, column=15, value=r.get("prompts_all_apps"))
            ws.cell(row=i, column=16, value=r.get("prompts_copilot_chat_work"))
            ws.cell(row=i, column=17, value=r.get("prompts_copilot_chat_web"))
            ws.cell(row=i, column=18, value=r.get("active_usage_days_all_apps"))
            ws.cell(row=i, column=19, value=r.get("last_activity_copilot_chat_work", ""))
            ws.cell(row=i, column=20, value=r.get("last_activity_copilot_chat_web", ""))
            ws.cell(row=i, column=21, value=r.get("last_activity_teams_copilot", ""))
            ws.cell(row=i, column=22, value=r.get("last_activity_word_copilot", ""))
            ws.cell(row=i, column=23, value=r.get("last_activity_excel_copilot", ""))
            ws.cell(row=i, column=24, value=r.get("last_activity_powerpoint_copilot", ""))
            ws.cell(row=i, column=25, value=r.get("last_activity_outlook_copilot", ""))
            ws.cell(row=i, column=26, value=r.get("last_activity_onenote_copilot", ""))
            ws.cell(row=i, column=27, value=r.get("last_activity_loop_copilot", ""))
            ws.cell(row=i, column=28, value=r.get("last_activity_m365_copilot_app", ""))
            ws.cell(row=i, column=29, value=r.get("last_activity_edge", ""))
            apply_row_style(ws, i, len(HEADERS))

    autofit_columns(ws, HEADERS)
