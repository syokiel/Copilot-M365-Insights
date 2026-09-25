"""KPI Trends — weekly/monthly series computed from the dated history in the
store (Viva Adoption, Copilot Studio sessions, credits, M365 service usage),
each with a line chart beside it. Unlike KPI History this needs no prior
syncs: a single import of a report with history already yields a trend."""
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.writers._style import apply_row_style, write_headers

# (title, note, rows key, [(header, field, format)], [(chart title, [fields])])
_SECTIONS = [
    ("M365 Copilot — weekly (Viva Copilot Adoption report)",
     "Enabled users = people whose Adoption-report span covers the week.",
     "copilot_weekly",
     [("Week", "week", None), ("Enabled Users", "enabled_users", "#,##0"),
      ("Active Users", "active_users", "#,##0"), ("Copilot Actions", "total_actions", "#,##0"),
      ("Chat", "chat_actions", "#,##0"), ("Teams", "teams_actions", "#,##0"),
      ("Outlook", "outlook_actions", "#,##0"), ("Word", "word_actions", "#,##0"),
      ("Excel", "excel_actions", "#,##0"), ("PowerPoint", "powerpoint_actions", "#,##0")],
     [("Copilot users per week", ["enabled_users", "active_users"]),
      ("Copilot actions per week by app", ["chat_actions", "teams_actions", "outlook_actions",
                                           "word_actions", "excel_actions", "powerpoint_actions"])]),
    ("Copilot Studio agents — monthly (Viva Copilot Studio report)",
     "Rates are % of sessions; CSAT is the average 1–5 score; Peak WAU sums users across agents.",
     "agent_sessions_monthly",
     [("Month", "month", None), ("Agents", "agents", "#,##0"), ("Sessions", "sessions", "#,##0"),
      ("Engaged", "engaged", "#,##0"), ("Resolution %", "resolution_rate", "0.0"),
      ("Escalation %", "escalation_rate", "0.0"), ("Abandon %", "abandon_rate", "0.0"),
      ("Avg CSAT", "csat_avg", "0.00"), ("Peak WAU", "peak_wau", "#,##0"),
      ("Autonomous Runs", "autonomous_runs", "#,##0"),
      ("Autonomous Success %", "autonomous_success_rate", "0.0")],
     [("Agent sessions & weekly users", ["sessions", "engaged", "peak_wau"]),
      ("Session outcomes (%)", ["resolution_rate", "escalation_rate", "abandon_rate"])]),
    ("Credits — monthly",
     "Capacity = Power Platform Copilot Studio capacity consumption; Consumption = Viva Consumption dashboard.",
     "credits_monthly",
     [("Month", "month", None), ("Capacity (billable)", "capacity_billable", "#,##0.0"),
      ("Capacity (non-billable)", "capacity_non_billable", "#,##0.0"),
      ("Consumption Credits", "consumption_credits", "#,##0.0"),
      ("Consumption People", "consumption_people", "#,##0")],
     [("Credits consumed per month", ["capacity_billable", "capacity_non_billable", "consumption_credits"])]),
    ("M365 services — monthly peak daily active users",
     "From the Office 365 active user counts export.",
     "services_monthly",
     [("Month", "month", None), ("Office 365", "office365", "#,##0"), ("Exchange", "exchange", "#,##0"),
      ("Teams", "teams", "#,##0"), ("SharePoint", "sharepoint", "#,##0"), ("OneDrive", "onedrive", "#,##0")],
     [("Peak daily active users", ["office365", "exchange", "teams", "sharepoint", "onedrive"])]),
]

_TITLE_FONT = Font(bold=True, size=12, color="1F4E79")
_NOTE_FONT = Font(italic=True, size=9, color="595959")
_CHART_ROWS = 16   # rows a 7.5 cm chart occupies


def write(ws: Worksheet, trends: dict[str, list[dict]]) -> None:
    row = 1
    widths: dict[int, int] = {}
    for title, note, key, cols, charts in _SECTIONS:
        data = trends.get(key) or []
        if not data:
            continue
        ws.cell(row=row, column=1, value=title).font = _TITLE_FONT
        ws.cell(row=row + 1, column=1, value=note).font = _NOTE_FONT
        header_row = row + 2
        write_headers(ws, [c[0] for c in cols], start_row=header_row)
        for i, rec in enumerate(data, start=header_row + 1):
            for j, (_header, field, fmt) in enumerate(cols, start=1):
                cell = ws.cell(row=i, column=j, value=rec.get(field))
                if fmt:
                    cell.number_format = fmt
            apply_row_style(ws, i, len(cols))
        last_row = header_row + len(data)
        for j, (header, _field, _fmt) in enumerate(cols, start=1):
            widths[j] = max(widths.get(j, 0), len(header) + 3, 12)

        # Charts beside the table, stacked if there's more than one.
        chart_col = get_column_letter(len(cols) + 2)
        chart_row = header_row
        if len(data) >= 2:
            fields = [c[1] for c in cols]
            categories = Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row)
            for chart_title, series in charts:
                chart = LineChart()
                chart.title = chart_title
                chart.height, chart.width = 7.5, 16
                for field in series:
                    c = fields.index(field) + 1
                    chart.add_data(Reference(ws, min_col=c, min_row=header_row, max_row=last_row),
                                   titles_from_data=True)
                chart.set_categories(categories)
                ws.add_chart(chart, f"{chart_col}{chart_row}")
                chart_row += _CHART_ROWS
        row = max(last_row, chart_row - 1) + 3

    if row == 1:
        ws.cell(row=1, column=1, value="— No dated history imported yet (Viva, Copilot Studio, credits or service counts) —")
    for j, w in widths.items():
        ws.column_dimensions[get_column_letter(j)].width = w
